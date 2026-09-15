"""Builds a small synthetic resume DOCX for offline tests.

Deliberately fictional (no real personal data) but structurally identical
to the real resumes this importer was built against: same paragraph
styles (Title/Subtitle/Heading 1/Role Header/List Bullet/Normal), same
section names, same overlapping-employment pattern (used to test
timeline.merge_intervals doesn't double-count).
"""

from pathlib import Path

import docx
from docx.enum.style import WD_STYLE_TYPE


def _add_role_header_style(document: docx.Document) -> None:
    styles = document.styles
    if "Role Header" not in [s.name for s in styles]:
        styles.add_style("Role Header", WD_STYLE_TYPE.PARAGRAPH)


def build(path: Path) -> None:
    document = docx.Document()
    _add_role_header_style(document)

    document.add_paragraph("Jane Example", style="Title")
    document.add_paragraph("Platform Engineer | Solutions Architect", style="Subtitle")
    document.add_paragraph("Testville, TX | jane@example.com")

    document.add_paragraph("Professional Summary", style="Heading 1")
    document.add_paragraph(
        "Fictional platform engineer with experience in architecture, cloud infrastructure, "
        "and customer-facing delivery, used only for offline importer tests."
    )

    document.add_paragraph("Core Capabilities", style="Heading 1")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Architecture"
    table.rows[0].cells[1].text = "Solution architecture, distributed systems"
    table.rows[1].cells[0].text = "Cloud"
    table.rows[1].cells[1].text = "AWS, Docker, CI/CD"

    document.add_paragraph("Professional Experience", style="Heading 1")

    p = document.add_paragraph(style="Role Header")
    p.add_run("Acme Corp  |  Platform Engineer\nTestville, TX  |  June 2020 - Present")
    document.add_paragraph(
        "Led cloud infrastructure automation and integration work across internal platforms.",
        style="List Bullet",
    )
    document.add_paragraph(
        "Built REST API integrations connecting internal services.", style="List Bullet"
    )

    # Deliberately overlaps with Acme Corp (2020-Present) for 2 years, to
    # exercise merge_intervals not double-counting.
    p = document.add_paragraph(style="Role Header")
    p.add_run("Side Consulting LLC  |  Technical Consultant\nRemote  |  June 2020 - June 2022")
    document.add_paragraph(
        "Advised clients on AWS cloud migrations as an independent consultant.",
        style="List Bullet",
    )

    p = document.add_paragraph(style="Role Header")
    p.add_run("Old Job Inc  |  Junior Developer\nTestville, TX  |  January 2016 - May 2020")
    document.add_paragraph("Built internal tools using PHP and MySQL.", style="List Bullet")

    document.add_paragraph("Selected Architecture Work", style="Heading 1")
    p = document.add_paragraph(style="Role Header")
    p.add_run("Internal Platform Rebuild  |  Architecture Lead\nAcme Ecosystem  |  Selected Initiative")
    document.add_paragraph(
        "Designed the caching, API, and deployment architecture for an internal platform.",
        style="List Bullet",
    )

    document.add_paragraph("Technical Foundation", style="Heading 1")
    document.add_paragraph("PHP  |  AWS  |  Docker  |  REST APIs")

    document.save(path)


if __name__ == "__main__":
    build(Path(__file__).parent / "sample_resume.docx")
